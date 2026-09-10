"""
Streamlit Frontend: Legal Document Assistant
"""

import os
import requests
import streamlit as st
import re
from io import BytesIO
from docx import Document

# Page Configuration
st.set_page_config(
    page_title="Legal AI Assistant",
    page_icon="assets/logo.svg",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- CUSTOM CSS: Toolbar ---
st.markdown(
    """
    <style>
        [data-testid="stHeadingWithActionElements"] h1 {
            padding-top: 0 !important;
            margin-top: 0 !important;
        }

        /* Force the actual content inside the tab to push itself down */
        div[role="tabpanel"] {
            padding-top: 1.5rem !important;
            margin-top: 1rem !important;
        }

        /* Hide the Deploy button */
        [data-testid="stAppDeployButton"] {
            display: none !important;
        }

        /* Hide the Running Man (Status Widget) */
        [data-testid="stStatusWidget"] {
            display: none !important;
        }
        
        /* Kill the invisible padding above the main Title / Logo */
        [data-testid="stHeadingWithActionElements"] {
            margin-top: -1.5rem !important;
            padding-top: 0 !important;
        }

        /* Match Streamlit action buttons to the height of status banners. */
        [data-testid="stButton"] button {
            min-height: 58px !important;
        }
    </style>
    """,
    unsafe_allow_html=True
)

# --- CUSTOM CSS: DROPBOX UPLOADER ---
st.markdown(
    """
    <style>
        /* The Main Outer Box */
        [data-testid="stFileUploaderDropzone"] {
            width: 100% !important;
            min-height: 520px !important; 
            border-radius: 24px !important;
            background-color: rgba(0, 97, 255, 0.04) !important;
            border: 2px dashed rgba(0, 97, 255, 0.4) !important;
            position: relative !important; 
            overflow: hidden !important; /* Keeps everything cleanly inside */
            transition: all 0.3s ease-in-out !important;
        }
        
        [data-testid="stFileUploaderDropzone"]:hover {
            background-color: rgba(0, 97, 255, 0.08) !important;
            border: 2px solid #0061FF !important;
        }

        /* Switch to a centered file preview after an upload. */
        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) {
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"])::before,
        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stFileUploaderDropzoneInstructions"],
        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stBaseButton-borderlessIcon"] {
            display: none !important;
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) > div {
            position: relative !important;
            width: 100% !important;
            height: 100% !important;
            display: block !important;
        }

        [data-testid="stFileUploaderDropzone"] [data-testid="stFileChips"] {
            position: absolute !important;
            top: 50% !important;
            left: 50% !important;
            transform: translate(-50%, -50%) !important;
            margin: 0 !important;
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stFileChip"] > div:first-child {
            width: 32px !important;
            height: 32px !important;
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stFileChip"] > div:first-child svg {
            display: none !important;
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stFileChip"] > div:first-child::after {
            content: "";
            width: 22px;
            height: 22px;
            border: 3px solid rgba(0, 97, 255, 0.2);
            border-top-color: #0061FF;
            border-radius: 50%;
            animation: legalUploaderSpin 0.8s linear infinite;
        }

        @keyframes legalUploaderSpin {
            to { transform: rotate(360deg); }
        }

        [data-testid="stFileUploaderDropzone"]:has([data-testid="stFileChips"]) [data-testid="stFileChip"] svg {
            display: block !important;
        }

        /* Pin Logo dead center using the pseudo-element */
        [data-testid="stFileUploaderDropzone"]::before {
            content: "";
            position: absolute !important;
            top: 40% !important; 
            left: 50% !important;
            transform: translate(-50%, -50%) !important; 
            
            width: 150px;  
            height: 150px;
            background-image: url("https://upload.wikimedia.org/wikipedia/commons/7/78/Dropbox_Icon.svg");
            background-size: contain;
            background-repeat: no-repeat;
            pointer-events: none; /* Allows dragging files right over the logo */
            z-index: 1 !important;
        }

        /* Hide default Streamlit cloud icon and "Drag and drop" text */
        [data-testid="stFileUploaderDropzone"] svg { display: none !important; }
        [data-testid="stFileUploaderDropzone"] > div > span { display: none !important; }

        /* Keep the controls together at the bottom of the dropzone */
        [data-testid="stFileUploaderDropzone"] {
            display: block !important;
            padding: 0 !important;
        }

        [data-testid="stFileUploaderDropzone"] > div {
            position: static !important;
            width: auto !important;
            height: auto !important;
            display: block !important;
            padding: 0 !important;
            z-index: 5 !important;
        }

        [data-testid="stFileUploaderDropzone"] > [data-testid="stFileUploaderDropzoneInstructions"] {
            position: absolute !important;
            right: 0 !important;
            bottom: 30px !important;
            left: 0 !important;
            display: flex !important;
            justify-content: center !important;
            align-items: center !important;
        }
        
        /* Clean up the Button */
        [data-testid="stFileUploaderDropzone"]:not(:has([data-testid="stFileChips"])) button {
            position: absolute !important;
            bottom: 62px !important;
            left: 50% !important;
            transform: translateX(-50%) !important;
            z-index: 6 !important;
            background-color: #0061FF !important;
            color: #ffffff !important;
            border: none !important;
            border-radius: 8px !important;
            padding: 0.8rem 2rem !important;
            font-weight: 600 !important;
            font-size: 18px !important;
            margin: 0 !important;
            margin-bottom: 12px !important; /* Space between button and 200mb text */
        }
        
        [data-testid="stFileUploaderDropzone"]:not(:has([data-testid="stFileChips"])) button:hover {
            background-color: #0050d1 !important;
        }

        /* Clean up the 200MB text */
        [data-testid="stFileUploaderDropzone"] [data-testid="stFileUploaderDropzoneInstructions"] span {
            display: block !important;
            position: static !important;
            width: auto !important;
            max-width: 100% !important;
            font-size: 15px !important; 
            color: #6b7280 !important;  
            margin: 0 !important;
            text-align: center !important;
        }

        /* Drop contract here label */
        [data-testid="stFileUploader"] label p {
            font-size: 18px !important;
            font-weight: 600 !important;
            margin-bottom: 12px !important;
        }
    </style>
    """,
    unsafe_allow_html=True
)

# Remove Streamlit's default massive top padding safely
st.markdown("""
    <style>
        .block-container {
            padding-top: 2rem !important;
            padding-bottom: 0rem !important;
        }
    </style>
""", unsafe_allow_html=True)

# Session State Initialization
if "analysis_result" not in st.session_state:
    st.session_state.analysis_result = None
if "total_tokens" not in st.session_state:
    st.session_state.total_tokens = 0
if "total_cost" not in st.session_state:
    st.session_state.total_cost = 0.0
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "chat_citations" not in st.session_state:
    st.session_state.chat_citations = []
if "workspace_uploader_version" not in st.session_state:
    st.session_state.workspace_uploader_version = 0

# 3. Environment & Directory Setup
BACKEND_URL = os.getenv("BACKEND_URL", "http://localhost:8000/analyze")
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.abspath(os.path.join(BASE_DIR, "..", "data", "sample_contracts"))

# ==============================================================================
# SIDEBAR
# ==============================================================================
with st.sidebar:
    st.title("Settings")
    
    st.subheader("AI Engine")
    selected_model = st.selectbox(
        "Model Tier",
        options=["gemini-3.6-flash", "gemini-3.8-flash"]
    )
    
    st.divider()
    
    # Create an empty placeholder. We will inject the metrics here at the very end of the script!
    usage_placeholder = st.empty()
    
    st.divider()
    st.caption("Legal Assistant Capstone v2.0")


def get_contract_files():
    if not os.path.exists(DATA_DIR):
        return []
    return sorted([f for f in os.listdir(DATA_DIR) if f.endswith(".txt")])


def load_contract_text(filename):
    filepath = os.path.join(DATA_DIR, filename)
    if not os.path.exists(filepath):
        return ""
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def highlight_text(full_text, flagged_clauses):
    highlighted = full_text
    
    for clause in flagged_clauses:
        snippet = clause.get("original_text", "")
        if not snippet:
            continue
            
        risk = clause.get("risk_level", "HIGH")
        bg_color = "#fee2e2" if risk == "HIGH" else ("#eff6ff" if risk == "INFO" else "#fef9c3")
        text_color = "#991b1b" if risk == "HIGH" else ("#1e3a8a" if risk == "INFO" else "#854d0e")
        border = "1px solid #f87171" if risk == "HIGH" else ("1px solid #60a5fa" if risk == "INFO" else "1px solid #facc15")
        reason_escaped = clause.get("reason", "").replace('"', '&quot;')
        
        # Clean the AI's string and escape special regex characters
        pattern = re.escape(snippet.strip())
        
        # Make the search flexible: treat any space or newline as "one or more whitespaces"
        pattern = pattern.replace(r'\ ', r'\s+').replace(r'\n', r'\s+')
        
        # \g<0> tells regex to insert the EXACT text it found in the document, preserving original formatting
        html_tag = (
            f'<mark style="background-color: {bg_color}; color: {text_color}; '
            f'border: {border}; font-weight: bold; padding: 2px 4px; '
            rf'border-radius: 4px;" title="{reason_escaped}">\g<0></mark>'
        )
        
        try:
            # Search and replace flexibly, ignoring case sensitivity
            highlighted = re.sub(pattern, html_tag, highlighted, flags=re.IGNORECASE)
        except Exception:
            # If regex fails for some weird edge case, fallback to exact string matching
            highlighted = highlighted.replace(snippet, html_tag)

    return highlighted


def generate_word_report(analysis_data, document_name, mode="Compliance Check"):
    """Generates a professional Word document from the AI analysis results in memory."""
    doc = Document()
    
    # Title and Metadata
    doc.add_heading('Legal Analysis Audit Report', 0)
    doc.add_paragraph(f"Document Analyzed: {document_name}")
    doc.add_paragraph(f"Analysis Mode: {mode}")
    
    # Executive Summary
    doc.add_heading('Executive Summary', level=1)
    doc.add_paragraph(analysis_data.get('summary', 'No summary provided.'))
    
    flagged = analysis_data.get("flagged_clauses", [])
    
    # Handle Compliance vs Abstraction
    if not flagged and mode == "Compliance Check":
         doc.add_paragraph("RESULT: Contract is fully compliant with the playbook. No high-risk clauses detected.")
    else:
        doc.add_heading('Detailed Findings', level=1)
        for clause in flagged:
            # Use different formatting based on risk level
            risk = clause.get('risk_level', 'INFO')
            
            doc.add_heading(f"{clause.get('clause_title', 'Clause')}", level=2)
            
            # Add bold labels for the metadata
            p_risk = doc.add_paragraph()
            p_risk.add_run("Risk Level: ").bold = True
            p_risk.add_run(risk)
            
            p_reason = doc.add_paragraph()
            p_reason.add_run("Details/Reason: ").bold = True
            p_reason.add_run(clause.get('reason', 'N/A'))
            
            if risk != "INFO":
                p_redline = doc.add_paragraph()
                p_redline.add_run("Proposed Redline:\n").bold = True
                p_redline.add_run(clause.get('proposed_redline', 'N/A'))
    
    # Save to a byte stream (no files created on disk!)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return buffer


# ==============================================================================
# MAIN WORKSPACE
# ==============================================================================
st.title("Interactive Legal Document Assistant")

tab_demo, tab_upload, tab_history = st.tabs([
    "Demo Gallery", 
    "Custom Workspace", 
    "History"
])

# ------------------------------------------------------------------------------
# TAB 1: DEMO GALLERY
# ------------------------------------------------------------------------------
with tab_demo:
    available_contracts = get_contract_files()
    col1, col2 = st.columns([6, 4])
    
    with col2:
        st.subheader("Rule Configuration")
        
        # Native Streamlit container to prevent page scrolling
        with st.container(height=640):
            selected_file = st.selectbox(
                "Select a Contract:",
                options=available_contracts if available_contracts else ["No contracts found"],
                key="demo_select"
            )
            
            playbook_rule = st.text_area(
                "Legal Playbook Rule:",
                value="The governing law of the contract must be the State of Delaware.",
                key="demo_rule",
                height=100
            )
            
            analyze_button = st.button(
                "Analyze Contract", 
                type="primary", 
                use_container_width=True, 
                key="demo_btn"
            )
            
            raw_text = ""
            if available_contracts and selected_file in available_contracts:
                raw_text = load_contract_text(selected_file)
            
            if analyze_button:
                if not raw_text:
                    st.warning("Please select a valid contract file first.")
                else:
                    with st.spinner("Analyzing against legal playbook..."):
                        try:
                            payload = {
                                "filename": selected_file,
                                "contract_text": raw_text,
                                "playbook_rule": playbook_rule,
                                "model": selected_model
                            }
                            response = requests.post(BACKEND_URL, json=payload, timeout=180)
                            response.raise_for_status()
                            
                            data = response.json()
                            st.session_state.analysis_result = data
                            
                            st.session_state.total_tokens += data.get("input_tokens", 0) or 0
                            st.session_state.total_cost += data.get("estimated_cost_usd", 0.0) or 0.0
                            
                        except Exception as e:
                            st.error(f"Analysis failed: {e}")
            
            result = st.session_state.analysis_result
            if result:
                st.divider()
                if result.get("is_compliant"):
                    st.success("Contract is fully compliant with the playbook.")
                else:
                    st.error("Non-Compliant Clauses Detected")
                    
                st.info(f"Summary: {result.get('summary')}")
                
                flagged = result.get("flagged_clauses", [])
                if flagged:
                    st.markdown(f"**Identified Violations ({len(flagged)}):**")
                    for clause in flagged:
                        title = clause.get("clause_title", "Flagged Clause")
                        with st.expander(title, expanded=True):
                            st.markdown(f"**Risk Level:** `{clause.get('risk_level', 'HIGH')}`")
                            st.markdown(f"**Reason:** {clause.get('reason', 'N/A')}")
                            st.markdown("**Proposed Redline:**")
                            st.markdown(clause.get("proposed_redline", "N/A"))
    
    with col1:
        st.subheader("Document Viewer")
        
        with st.container(height=640):
            if raw_text:
                display_text = raw_text
                if st.session_state.analysis_result:
                    display_text = highlight_text(
                        raw_text, 
                        st.session_state.analysis_result.get("flagged_clauses", [])
                    )
                
                display_text_html = display_text.replace("\n", "<br>")
                
                # Hardcoded white background and black text for dark mode compatibility
                st.markdown(
                    f"""
                    <div style="
                        background-color: #ffffff; 
                        color: #000000; 
                        padding: 20px; 
                        border-radius: 5px; 
                        border: 1px solid #ccc; 
                        font-family: sans-serif; 
                        font-size: 14px; 
                        line-height: 1.6;
                    ">
                        {display_text_html}
                    </div>
                    """,
                    unsafe_allow_html=True
                )
            else:
                st.info("Select a contract from the panel on the right to preview it.")

# ------------------------------------------------------------------------------
# TAB 2: CUSTOM WORKSPACE
# ------------------------------------------------------------------------------
with tab_upload:
    # State Management for the Custom Workspace
    if "custom_file_name" not in st.session_state:
        st.session_state.custom_file_name = None
    if "custom_raw_text" not in st.session_state:
        st.session_state.custom_raw_text = None
    if "custom_analysis" not in st.session_state:
        st.session_state.custom_analysis = None

    def clear_workspace():
        try:
            cleanup_url = BACKEND_URL.replace("/analyze", "/cleanup")
            payload = {"filename": st.session_state.custom_file_name}
            cleanup_response = requests.post(
                cleanup_url,
                json=payload,
                timeout=10
            )
            cleanup_response.raise_for_status()
        except Exception as e:
            st.error(f"Workspace cleanup failed: {e}")
            return

        st.session_state.custom_file_name = None
        st.session_state.custom_raw_text = None
        st.session_state.custom_analysis = None
        st.session_state.chat_history = []
        st.session_state.chat_citations = []
        st.session_state.workspace_uploader_version += 1

    col1, col2 = st.columns([6, 4])
    
    # ---------------- RIGHT COLUMN: RULES & EXPORT ----------------
    with col2:
        st.subheader("Custom Analysis")
        with st.container(height=640):
            # Only show settings if a file is uploaded
            if not st.session_state.custom_raw_text:
                st.info("Upload a document on the left to unlock analysis settings.")
            else:
                def reset_analysis_state():
                    st.session_state.custom_analysis = None
                    st.session_state.chat_citations = [] # Wipes highlights when changing modes
                
                analysis_mode = st.radio(
                    "Analysis Mode:", 
                    ["Compliance Check", "General Abstraction", "Q&A Chat"], 
                    horizontal=True,
                    on_change=reset_analysis_state,
                    key="mode_toggle"
                )

                if analysis_mode != "Q&A Chat":
                    if analysis_mode == "Compliance Check":
                        playbook_rule = st.text_area(
                            "Legal Playbook Rule:",
                            value="The governing law must be the State of Delaware.",
                            key="custom_rule",
                            height=90
                        )
                        mode_param = "compliance"
                    else:
                        st.info("The AI will extract core rights, obligations, financials, and termination terms without a specific rule.")
                        playbook_rule = "General Contract Abstraction"
                        mode_param = "abstraction"
                    
                    analyze_btn = st.button("Analyze Document", type="primary", use_container_width=True)
                    
                    if analyze_btn:
                        with st.spinner("Analyzing document..."):
                            try:
                                payload = {
                                    "filename": st.session_state.custom_file_name,
                                    "contract_text": st.session_state.custom_raw_text,
                                    "playbook_rule": playbook_rule,
                                    "model": selected_model,
                                    "mode": mode_param
                                }
                                response = requests.post(BACKEND_URL, json=payload, timeout=180)
                                response.raise_for_status()
                                data = response.json()
                                st.session_state.custom_analysis = data
                                
                                st.session_state.total_tokens += data.get("input_tokens", 0) or 0
                                st.session_state.total_cost += data.get("estimated_cost_usd", 0.0) or 0.0
                            except Exception as e:
                                st.error(f"Analysis failed: {e}")

                else:
                    st.info("Ask open-ended questions about your contract. The AI will retrieve the answers directly from the text.")
                    
                    # Display existing chat history
                    for msg in st.session_state.chat_history:
                        with st.chat_message(msg["role"]):
                            st.markdown(msg["content"])
                    
                    # Chat Input Box
                    if prompt := st.chat_input("E.g., What are the payment terms in this contract?"):
                        # Append user message to UI
                        st.session_state.chat_history.append({"role": "user", "content": prompt})
                        with st.chat_message("user"):
                            st.markdown(prompt)
                            
                        # Fetch AI response
                        with st.chat_message("assistant"):
                            with st.spinner("Searching document..."):
                                try:
                                    chat_url = BACKEND_URL.replace("/analyze", "/chat")
                                    payload = {
                                        "filename": st.session_state.custom_file_name,
                                        "query": prompt,
                                        "model": selected_model
                                    }
                                    res = requests.post(chat_url, json=payload, timeout=180)
                                    res.raise_for_status()
                                    
                                    data = res.json()
                                    answer = data.get("answer", "No answer received.")
                                    quotes = data.get("exact_quotes", [])
                                    
                                    # Save to memory
                                    st.session_state.chat_history.append({"role": "assistant", "content": answer})
                                    st.session_state.chat_citations = quotes

                                    # Force a UI refresh so the left column immediately highlights the quotes
                                    st.rerun()

                                except Exception as e:
                                    st.error(f"Chat failed: {e}")
                
                # Render Results
                result = st.session_state.custom_analysis
                if result:
                    st.divider()
                    if analysis_mode == "Compliance Check":
                        if result.get("is_compliant"):
                            st.success("Contract is fully compliant with the playbook.")
                        else:
                            st.error("Non-Compliant Clauses Detected")
                    else:
                        st.success("General Abstraction Complete")
                        
                    st.info(f"**Summary:** {result.get('summary')}")
                    
                    flagged = result.get("flagged_clauses", [])
                    for clause in flagged:
                        with st.expander(f"{clause.get('clause_title', 'Clause')}", expanded=True):
                            st.markdown(f"**Level:** `{clause.get('risk_level', 'HIGH')}`")
                            st.markdown(f"**Details:** {clause.get('reason', 'N/A')}")
                            if clause.get('risk_level') != "INFO":
                                st.markdown(f"**Proposed Redline:**\n{clause.get('proposed_redline', 'N/A')}")

                    # --- EXPORT TO AUDIT REPORT ---
                    st.divider()
                    report_text = f"LEGAL COMPLIANCE AUDIT\nDocument: {st.session_state.custom_file_name}\nModel Used: {selected_model}\n"
                    report_text += f"{'='*50}\n\nOVERALL SUMMARY:\n{result.get('summary')}\n\n{'='*50}\n\n"
                    
                    if not flagged:
                        report_text += "RESULT: Contract is fully compliant.\n"
                    else:
                        report_text += f"VIOLATIONS FOUND ({len(flagged)}):\n\n"
                        for c in flagged:
                            report_text += f"CLAUSE: {c.get('clause_title')}\nRISK: {c.get('risk_level')}\n"
                            report_text += f"ISSUE: {c.get('reason')}\nREDLINE: {c.get('proposed_redline')}\n"
                            report_text += "-"*30 + "\n"
                            
                    # Generate the Word Document in memory
                    docx_buffer = generate_word_report(result, st.session_state.custom_file_name, analysis_mode)
                                        
                    st.download_button(
                        label="Download Word Report (.docx)",
                        data=docx_buffer,
                        file_name=f"Audit_{st.session_state.custom_file_name}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )

    # ---------------- LEFT COLUMN: UPLOADER OR VIEWER ----------------
    with col1:
        st.subheader("Document Workspace")
        
        with st.container(height=640):
            # STATE 1: Empty - Show Drag and Drop
            if not st.session_state.custom_raw_text:
                uploaded_file = st.file_uploader(
                    "Drop contract here",
                    type=["txt", "pdf"],
                    key=f"workspace_uploader_{st.session_state.workspace_uploader_version}"
                )
                
                if uploaded_file:
                    with st.spinner(f"Extracting and Vectorizing `{uploaded_file.name}`..."):
                        try:
                            # Send file to new FastAPI upload route
                            upload_url = BACKEND_URL.replace("/analyze", "/upload")
                            files = {"file": (uploaded_file.name, uploaded_file.getvalue(), uploaded_file.type)}
                            res = requests.post(upload_url, files=files)
                            res.raise_for_status()
                            
                            upload_data = res.json()
                            # Transition to State 2
                            st.session_state.custom_file_name = upload_data["filename"]
                            st.session_state.custom_raw_text = upload_data["text"]
                            st.session_state.custom_analysis = None
                            st.rerun() # Force UI to update immediately
                        except Exception as e:
                            st.error(f"Upload failed: {e}")
                else:
                    st.write("Upload a `.txt` or text-based `.pdf` contract. The backend will instantly extract, chunk, and vectorize it.")
                            
            # STATE 2: Uploaded - Show Document Viewer
            else:
                top_col1, top_col2 = st.columns([7, 3])
                with top_col1:
                    st.success(f"Active: `{st.session_state.custom_file_name}`")
                with top_col2:
                    st.button(
                        "Clear Workspace",
                        use_container_width=True,
                        on_click=clear_workspace
                    )
                
                st.divider()

                st.subheader("Document Viewer")
                
                if st.session_state.custom_raw_text:
                    import re
                    display_text = st.session_state.custom_raw_text
                    
                    # Safely fetch the mode from global session state
                    current_mode = st.session_state.get("mode_toggle", "Compliance Check")
                    
                    # --- Q&A Chat Highlighting ---
                    if current_mode == "Q&A Chat":
                        if st.session_state.get("chat_citations"):
                            for quote in st.session_state.chat_citations:
                                # Break the quote into individual words
                                words = quote.strip().split()
                                
                                # Only highlight if it's a meaningful phrase (more than 3 words)
                                if len(words) > 3: 
                                    # 3. Create a flexible pattern that ignores spacing/newlines between words (\s+)
                                    flexible_pattern = r'\s+'.join(re.escape(w) for w in words)
                                    
                                    html_tag = (
                                        f'<mark style="background-color: #f3e8ff; color: #6b21a8; '
                                        f'border: 1px solid #c084fc; font-weight: bold; padding: 2px 4px; '
                                        rf'border-radius: 4px;" title="Source Citation">\g<0></mark>'
                                    )
                                    # Apply the flexible pattern
                                    display_text = re.sub(f'({flexible_pattern})', html_tag, display_text, flags=re.IGNORECASE)
                    
                    # --- Compliance/Abstraction Highlighting ---
                    elif st.session_state.get("custom_analysis"):
                        display_text = highlight_text(
                            display_text, 
                            st.session_state.custom_analysis.get("flagged_clauses", [])
                        )
                    
                    # ---  HTML AND CSS ---
                    display_text_html = display_text.replace("\n", "<br>")
                    
                    st.markdown(
                        f"""
                        <div style="
                            background-color: #ffffff; 
                            color: #1f2937; 
                            padding: 24px; 
                            border-radius: 6px; 
                            border: 1px solid #d1d5db; 
                            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; 
                            font-size: 13.5px; 
                            line-height: 1.7;
                            box-shadow: 0 1px 3px rgba(0,0,0,0.06);
                        ">
                            {display_text_html}
                        </div>
                        """,
                        unsafe_allow_html=True
                    )
                else:
                    st.info("Upload a document to view it here.")


# ------------------------------------------------------------------------------
# TAB 3: HISTORY DASHBOARD
# ------------------------------------------------------------------------------
with tab_history:
    def delete_record_callback(record_id):
        try:
            delete_url = BACKEND_URL.replace("/analyze", f"/history/{record_id}")
            requests.delete(delete_url, timeout=5)
        except Exception as e:
            st.toast(f"Failed to delete record: {e}")

    st.subheader("Analysis Audit Trail")
    
    # Fetch data from the FastAPI backend
    history_data = []
    try:
        history_url = BACKEND_URL.replace("/analyze", "/history")
        res = requests.get(history_url, timeout=10)
        res.raise_for_status()
        history_data = res.json().get("history", [])
    except Exception as e:
        st.error(f"Could not connect to history database: {e}")

    # Render KPIs
    if not history_data:
        st.info("No analysis history found. Run an analysis in the Demo or Custom workspace to see it here.")
    else:
        total_runs = len(history_data)
        total_spend = sum(item.get("estimated_cost_usd", 0.0) for item in history_data)
        
        kpi1, kpi2, kpi3 = st.columns(3)
        kpi1.metric(label="Total Documents Analyzed", value=total_runs)
        kpi2.metric(label="Total API Spend", value=f"${total_spend:.6f}")
        kpi3.metric(label="Data Privacy Status", value="Secured locally")
        
        st.divider()

        # Render the Master-Detail List
        for record in history_data:
            status_badge = "Compliant" if record["is_compliant"] else "Issues Found"
            
            # Truncate the rule for the header if it's too long
            rule_preview = record['playbook_rule']
            if len(rule_preview) > 60:
                rule_preview = rule_preview[:60] + "..."
                
            header_title = f"{status_badge} | {record['timestamp']} | {record['filename']} | Rule: {rule_preview}"
            
            with st.expander(header_title):
                st.markdown(f"**Model:** `{record['model_used']}` | **Cost for this run:** `${record.get('estimated_cost_usd', 0.0):.6f}`")
                st.info(f"**Summary:** {record['summary']}")
                
                # Show flagged clauses if any
                flagged = record.get("flagged_clauses", [])
                if flagged:
                    for clause in flagged:
                        st.error(f"**{clause.get('clause_title', 'Flagged Clause')}**\n\n**Risk:** `{clause.get('risk_level')}`\n\n**Issue:** {clause.get('reason')}\n\n**Redline:**\n{clause.get('proposed_redline')}")
                
                st.divider()
                
                # Action Buttons (Download & Delete)
                action_col1, action_col2, action_col3 = st.columns([3, 3, 4])
                
                with action_col1:
                    # Determine the mode historically based on if there's a rule
                    hist_mode = "General Abstraction" if record['playbook_rule'] == "General Contract Abstraction" else "Compliance Check"
                    
                    # Generate the Word Document
                    docx_buffer = generate_word_report(record, record['filename'], hist_mode)
                    
                    st.download_button(
                        label="Download Word Report",
                        data=docx_buffer,
                        file_name=f"Audit_{record['filename']}_{record['timestamp'].replace(':', '-')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"dl_{record['id']}",
                        use_container_width=True
                    )
                    
                with action_col2:
                    st.button(
                        "Delete Record", 
                        key=f"del_{record['id']}", 
                        on_click=delete_record_callback,
                        args=(record['id'],),
                        use_container_width=True
                    )


# ==============================================================================
# INJECT SIDEBAR METRICS
# ==============================================================================
with usage_placeholder.container():
    st.subheader("Session Usage")
    st.metric(label="Tokens Processed", value=f"{st.session_state.total_tokens:,}")
    st.metric(label="Est. Enterprise Cost", value=f"${st.session_state.total_cost:.5f}")